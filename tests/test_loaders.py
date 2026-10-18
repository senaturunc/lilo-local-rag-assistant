import io
import pytest
from docx import Document
from loaders import extract_text


def test_txt_is_decoded():
    assert extract_text("a.txt", "Merhaba dünya".encode("utf-8")) == "Merhaba dünya"


def test_txt_with_bom_is_cleaned():
    assert extract_text("a.txt", "\ufeffMerhaba".encode("utf-8")) == "Merhaba"


def test_docx_paragraphs_are_extracted():
    doc = Document()
    doc.add_paragraph("Birinci paragraf.")
    doc.add_paragraph("İkinci paragraf.")
    buf = io.BytesIO()
    doc.save(buf)

    text = extract_text("a.docx", buf.getvalue())
    assert "Birinci paragraf." in text and "İkinci paragraf." in text


def test_unsupported_extension_raises():
    with pytest.raises(ValueError):
        extract_text("a.xyz", b"data")